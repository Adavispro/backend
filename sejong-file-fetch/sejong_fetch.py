"""Select Sejong files by source date, stage originals, and export dated MDB tables."""

import csv
import hashlib
import json
import os
import re
import shutil
import subprocess
import time
from datetime import date, datetime, timedelta, timezone
from pathlib import Path


DATE_IN_NAME = re.compile(r"(?<!\d)(\d{4}-\d{2}-\d{2})(?!\d)")
DATABASE_EXTENSIONS = {".mdb", ".accdb", ".fdb"}


def now_utc():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def timestamp():
    return datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")


def digest_file(path):
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def save_json(path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    temporary.write_text(json.dumps(payload, indent=2, ensure_ascii=False, default=str) + "\n", encoding="utf-8")
    os.replace(temporary, path)


def classify(path):
    if path.suffix.lower() in DATABASE_EXTENSIONS:
        return "database"
    name = path.stem.lower().replace("_", "").replace("-", "")
    if "productionreport" in name:
        return "production_report"
    if "alarm" in name:
        return "alarms"
    if "audit" in name:
        return "audit"
    if "loginout" in name or "loginlogout" in name:
        return "login_logout"
    if "operatinghistory" in name:
        return "operating_history"
    if "pressuredata" in name:
        return "pressure_data"
    if "tabletchecker" in name:
        return "tablet_checker"
    return "other_spreadsheet"


def selected_dates(config):
    mode = config.get("date_mode", "EXACT").upper()
    if mode == "TODAY":
        return [date.today()]
    if mode == "EXACT":
        return [date.fromisoformat(config["date"])]
    if mode == "RANGE":
        first = date.fromisoformat(config["start_date"])
        last = date.fromisoformat(config["end_date"])
        if last < first or (last - first).days > 30:
            raise ValueError("Date range must contain 1 to 31 days")
        return [first + timedelta(days=offset) for offset in range((last - first).days + 1)]
    raise ValueError("date_mode must be EXACT, RANGE, or TODAY")


def source_date(relative):
    folder_dates = []
    for part in relative.parts[:-1]:
        try:
            folder_dates.append(date.fromisoformat(part))
        except ValueError:
            continue
    match = DATE_IN_NAME.search(relative.name)
    filename_date = date.fromisoformat(match.group(1)) if match else None
    return (folder_dates[-1] if folder_dates else filename_date), filename_date


def csv_rows(path):
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def write_csv(path, rows, columns):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def report_fields(sheets):
    wanted = {
        "batchno": "batch_no", "productname": "product_name", "machinename": "machine_name",
        "userid": "user_id", "stationno": "station_no",
    }
    found = {}
    for sheet in sheets:
        by_row = {}
        for cell in sheet["cells"]:
            by_row.setdefault(cell["row"], []).append(cell)
        for cells in by_row.values():
            cells.sort(key=lambda item: item["column"])
            for index, cell in enumerate(cells):
                label = re.sub(r"[^a-z0-9]", "", str(cell["value"]).lower())
                field = wanted.get(label)
                if field and field not in found and index + 1 < len(cells):
                    found[field] = cells[index + 1]["value"]
    return found


def spreadsheet_cells(path, max_cells):
    sheets = []
    if path.suffix.lower() == ".xls":
        try:
            import xlrd
        except ImportError:
            raise RuntimeError("Install xlrd to extract .xls cells: python -m pip install xlrd") from None
        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            for sheet in book.sheets():
                cells = []
                for row in range(sheet.nrows):
                    for column in range(sheet.ncols):
                        cell = sheet.cell(row, column)
                        if cell.value in ("", None):
                            continue
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, book.datemode).isoformat()
                        cells.append({"row": row + 1, "column": column + 1, "value": value})
                        if len(cells) > max_cells:
                            raise RuntimeError(f"Spreadsheet exceeds max_cells_per_sheet={max_cells}")
                sheets.append({"name": sheet.name, "cells": cells})
        finally:
            book.release_resources()
    else:
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise RuntimeError("Install openpyxl to extract .xlsx/.xlsm cells: python -m pip install openpyxl") from None
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in book.worksheets:
                cells = []
                for row_number, row in enumerate(sheet.iter_rows(), start=1):
                    for column_number, cell in enumerate(row, start=1):
                        if cell.value in ("", None):
                            continue
                        value = cell.value.isoformat() if isinstance(cell.value, (date, datetime)) else cell.value
                        cells.append({"row": row_number, "column": column_number, "value": value})
                        if len(cells) > max_cells:
                            raise RuntimeError(f"Spreadsheet exceeds max_cells_per_sheet={max_cells}")
                sheets.append({"name": sheet.title, "cells": cells})
        finally:
            book.close()
    return sheets


class SejongFetcher:
    def __init__(self, config, root, logger):
        self.config = config
        self.root = root
        self.logger = logger
        self.source_root = Path(config["source_path"])
        staging = Path(config.get("staging_folder", "./staging"))
        self.staging_root = staging if staging.is_absolute() else root / staging
        self.index_path = self.staging_root / "metadata" / "file_index.json"
        self.extensions = {extension.lower() for extension in config["include_extensions"]}
        self.dates = selected_dates(config)
        self.selected = set(self.dates)
        self.index = {}

    def _candidate_files(self):
        if not self.source_root.is_dir():
            raise FileNotFoundError(f"Source folder is unavailable: {self.source_root}")
        for path in self.source_root.rglob("*"):
            if not path.is_file() or path.suffix.lower() not in self.extensions:
                continue
            if path.suffix.lower() in {".ldb", ".laccdb"} or path.name.startswith("~$"):
                continue
            relative = path.relative_to(self.source_root)
            file_date, filename_date = source_date(relative)
            if file_date is not None and file_date in self.selected:
                if filename_date is not None and filename_date != file_date:
                    self.logger.warning("source=%s filename_date=%s folder_date=%s", path, filename_date, file_date)
                yield path, relative, file_date
            elif (file_date is None and path.suffix.lower() in DATABASE_EXTENSIONS
                  and self.config.get("include_root_databases", True)):
                yield path, relative, None

    def _copy(self, source, relative, file_date, category):
        attempts = max(1, int(self.config.get("retry_count", 3)))
        delay = max(0, float(self.config.get("retry_delay_seconds", 5)))
        destination_root = self.staging_root / "raw" / (file_date.isoformat() if file_date else "databases") / category / relative.parent
        destination_root.mkdir(parents=True, exist_ok=True)
        last_error = None
        for attempt in range(attempts):
            temporary = destination_root / (".copy-" + timestamp() + "-" + source.name)
            try:
                before = source.stat()
                shutil.copy2(source, temporary)
                after = source.stat()
                if (before.st_mtime_ns, before.st_size) != (after.st_mtime_ns, after.st_size):
                    raise OSError("Source changed during copy")
                checksum = digest_file(temporary)
                target = destination_root / f"{source.stem}__{timestamp()}__{checksum[:8]}{source.suffix}"
                os.replace(temporary, target)
                return target, checksum, after
            except OSError as error:
                last_error = error
                temporary.unlink(missing_ok=True)
                if attempt + 1 < attempts:
                    time.sleep(delay)
        raise last_error

    def _export_table(self, database, table, output):
        if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_]*", table):
            raise ValueError("Unsafe Access table name")
        if os.name == "nt":
            command = ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File",
                       str(self.root / "mdb_export.ps1"), "-DatabasePath", str(database),
                       "-TableName", table, "-OutputPath", str(output)]
            result = subprocess.run(command, capture_output=True, text=True, timeout=180)
            if result.returncode == 3:
                return False
            if result.returncode != 0:
                raise RuntimeError(result.stderr.strip()[-500:] or "Access export failed")
            return True
        binary_dir = os.environ.get("MDBTOOLS_BIN_DIR")
        tables_command = str(Path(binary_dir) / "mdb-tables") if binary_dir else shutil.which("mdb-tables")
        export_command = str(Path(binary_dir) / "mdb-export") if binary_dir else shutil.which("mdb-export")
        if not tables_command or not export_command:
            raise RuntimeError("MDB Tools is needed to export Access tables on this computer")
        listing = subprocess.run([tables_command, str(database)], capture_output=True, text=True, timeout=60, check=True)
        if table not in listing.stdout.split():
            return False
        output.parent.mkdir(parents=True, exist_ok=True)
        with output.open("wb") as handle:
            result = subprocess.run([export_command, str(database), table], stdout=handle,
                                    stderr=subprocess.PIPE, timeout=180)
        if result.returncode != 0:
            output.unlink(missing_ok=True)
            raise RuntimeError("MDB table export failed")
        return True

    def _derive_events(self, folder):
        raw_events = csv_rows(folder / "sawc_data.csv")
        lookups = {(row.get("ClassA", ""), row.get("ClassB", "")): row.get("FullName", "")
                   for row in csv_rows(folder / "classb_lookup.csv")}
        events = []
        for row in raw_events:
            event = dict(row)
            event["EventName"] = lookups.get((row.get("ClassA", ""), row.get("ClassB", "")), "")
            events.append(event)
        columns = list(raw_events[0]) + ["EventName"] if raw_events else ["EventName"]
        write_csv(folder / "audit_events.csv", events, columns)
        alarms = [row for row in events if row.get("ClassA") == "0030" or "alarm" in row["EventName"].lower()]
        write_csv(folder / "alarms_and_safety.csv", alarms, columns)
        login = [row for row in events if row.get("ClassA") == "0000"]
        write_csv(folder / "login_logout.csv", login, columns)
        return {"audit_events": len(events), "alarms_and_safety": len(alarms), "login_logout": len(login)}

    def _export_database_date(self, database, source_key, record, target_date):
        date_key = target_date.isoformat()
        prior = record.setdefault("exports", {}).get(date_key, {})
        if prior.get("source_checksum") == record["checksum"] and prior.get("status") in {"SUCCESS", "NO_DATA"}:
            return
        folder = self.staging_root / "extracted" / date_key / (database.stem + "_" + hashlib.sha256(source_key.encode()).hexdigest()[:8]) / record["checksum"][:8]
        suffix = target_date.strftime("%Y%m%d")
        counts = {}
        files = []
        try:
            sawc = folder / "sawc_data.csv"
            awc = folder / "awc_data.csv"
            sawc_exists = self._export_table(database, "SAWC_DATA_" + suffix, sawc)
            awc_exists = self._export_table(database, "AWC_" + suffix, awc)
            if sawc_exists:
                files.append(str(sawc))
                if self._export_table(database, "ClassA", folder / "classa_lookup.csv"):
                    files.append(str(folder / "classa_lookup.csv"))
                if self._export_table(database, "ClassB", folder / "classb_lookup.csv"):
                    files.append(str(folder / "classb_lookup.csv"))
                counts.update(self._derive_events(folder))
                files.extend(str(folder / name) for name in ("audit_events.csv", "alarms_and_safety.csv", "login_logout.csv"))
            if awc_exists:
                files.append(str(awc))
                counts["awc_data"] = len(csv_rows(awc))
            status = "SUCCESS" if sawc_exists or awc_exists else "NO_DATA"
            record["exports"][date_key] = {
                "status": status, "source_checksum": record["checksum"],
                "export_timestamp": now_utc(), "files": files, "row_counts": counts,
            }
            self.logger.info("database=%s date=%s export=%s counts=%s", database.name, date_key, status, counts)
        except (OSError, RuntimeError, subprocess.SubprocessError) as error:
            record["exports"][date_key] = {
                "status": "FAILED", "source_checksum": record["checksum"],
                "export_timestamp": now_utc(), "error_message": str(error),
            }
            self.logger.error("database=%s date=%s export=FAILED error=%s", database.name, date_key, error)

    def _export_spreadsheet(self, staged_file, source_key, record, file_date):
        prior = record.get("spreadsheet_export", {})
        if prior.get("source_checksum") == record["checksum"] and prior.get("status") == "SUCCESS":
            return
        date_key = file_date.isoformat()
        folder = (self.staging_root / "extracted" / date_key / record["category"]
                  / (staged_file.stem + "_" + hashlib.sha256(source_key.encode()).hexdigest()[:8])
                  / record["checksum"][:8])
        try:
            sheets = spreadsheet_cells(staged_file, int(self.config.get("max_cells_per_sheet", 50000)))
            summary = report_fields(sheets)
            output = folder / "cells.json"
            save_json(output, {"source_file": record["source_file"], "source_date": date_key,
                               "category": record["category"], "summary": summary, "sheets": sheets})
            record["spreadsheet_export"] = {
                "status": "SUCCESS", "source_checksum": record["checksum"],
                "output_file": str(output), "summary": summary, "export_timestamp": now_utc(),
            }
            self.logger.info("spreadsheet=%s date=%s export=SUCCESS", record["source_filename"], date_key)
        except (OSError, RuntimeError, ValueError) as error:
            record["spreadsheet_export"] = {
                "status": "FAILED", "source_checksum": record["checksum"],
                "error_message": str(error), "export_timestamp": now_utc(),
            }
            self.logger.error("spreadsheet=%s date=%s export=FAILED error=%s", record["source_filename"], date_key, error)

    def _process(self, source, relative, file_date):
        key = str(source.resolve())
        previous = self.index.get(key, {})
        category = classify(source)
        stat = source.stat()
        unchanged = (previous.get("status") == "SUCCESS" and previous.get("source_mtime_ns") == stat.st_mtime_ns
                     and previous.get("source_size") == stat.st_size and Path(previous.get("staged_file", "")).is_file())
        if unchanged:
            record = previous
            self.logger.info("source=%s reason=ALREADY_STAGED", source)
        else:
            destination, checksum, copied_stat = self._copy(source, relative, file_date, category)
            record = {
                "source_file": str(source), "source_relative_path": str(relative), "source_filename": source.name,
                "source_date": file_date.isoformat() if file_date else None, "category": category,
                "source_mtime_ns": copied_stat.st_mtime_ns, "source_size": copied_stat.st_size,
                "file_modified_timestamp": datetime.fromtimestamp(copied_stat.st_mtime, timezone.utc).isoformat(),
                "fetch_timestamp": now_utc(), "staged_file": str(destination), "checksum": checksum,
                "status": "SUCCESS", "exports": previous.get("exports", {}),
                "spreadsheet_export": previous.get("spreadsheet_export", {}),
                "versions": previous.get("versions", []) + [{"staged_file": str(destination), "checksum": checksum, "fetch_timestamp": now_utc()}],
            }
            self.logger.info("source=%s date=%s category=%s status=STAGED", source, record["source_date"], category)
        if (source.suffix.lower() in {".mdb", ".accdb"} and self.config.get("extract_mdb_tables", True)):
            for target_date in ([file_date] if file_date else self.dates):
                self._export_database_date(Path(record["staged_file"]), key, record, target_date)
        if (file_date and source.suffix.lower() in {".xls", ".xlsx", ".xlsm"}
                and self.config.get("extract_spreadsheets", True)):
            self._export_spreadsheet(Path(record["staged_file"]), key, record, file_date)
        self.index[key] = record
        save_json(self.index_path, self.index)

    def run_cycle(self):
        self.logger.info("cycle start dates=%s source=%s", ",".join(value.isoformat() for value in self.dates), self.source_root)
        if self.index_path.exists():
            with self.index_path.open(encoding="utf-8") as handle:
                self.index = json.load(handle)
        try:
            candidates = list(self._candidate_files())
        except OSError as error:
            self.logger.error("discovery failed: %s", error)
            return
        available = []
        skipped = 0
        for item in candidates:
            try:
                source, _, file_date = item
                stat = source.stat()
                prior = self.index.get(str(source.resolve()), {})
                unchanged = (prior.get("status") == "SUCCESS"
                             and prior.get("source_mtime_ns") == stat.st_mtime_ns
                             and prior.get("source_size") == stat.st_size
                             and Path(prior.get("staged_file", "")).is_file())
                export_needed = False
                if (unchanged and source.suffix.lower() in {".mdb", ".accdb"}
                        and self.config.get("extract_mdb_tables", True)):
                    for target_date in ([file_date] if file_date else self.dates):
                        export = prior.get("exports", {}).get(target_date.isoformat(), {})
                        if export.get("source_checksum") != prior.get("checksum") or export.get("status") not in {"SUCCESS", "NO_DATA"}:
                            export_needed = True
                            break
                if (unchanged and file_date and source.suffix.lower() in {".xls", ".xlsx", ".xlsm"}
                        and self.config.get("extract_spreadsheets", True)):
                    export = prior.get("spreadsheet_export", {})
                    if export.get("source_checksum") != prior.get("checksum") or export.get("status") != "SUCCESS":
                        export_needed = True
                if unchanged and not export_needed:
                    skipped += 1
                    continue
                available.append((stat.st_mtime_ns, item))
            except OSError as error:
                self.logger.error("source=%s status=FAILED error=%s", item[0], error)
        available.sort(key=lambda item: item[0], reverse=True)
        candidates = [item for _, item in available]
        limit = max(0, int(self.config.get("max_files_per_cycle", 500)))
        self.logger.info("files_found=%d already_staged=%d pending=%d files_selected=%d",
                         len(candidates) + skipped, skipped, len(candidates), min(len(candidates), limit))
        for source, relative, file_date in candidates[:limit]:
            try:
                self._process(source, relative, file_date)
            except (OSError, ValueError, RuntimeError) as error:
                self.logger.error("source=%s status=FAILED error=%s", source, error)
        self.logger.info("cycle end")
